from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


MAX_DOCX_CONTENT_CHARS = 200_000

DISCLAIMER = (
    "Tài liệu này được tạo tự động bởi hệ thống AI nhằm hỗ trợ phân tích và "
    "tham khảo. Nội dung không phải là bản án, quyết định, kết luận chính thức "
    "hoặc văn bản do cơ quan tư pháp ban hành."
)


def build_legal_analysis_docx(content: str) -> tuple[BytesIO, str]:
    content = content.strip()

    if not content:
        raise ValueError("Nội dung báo cáo không được để trống.")

    if len(content) > MAX_DOCX_CONTENT_CHARS:
        raise ValueError(
            f"Nội dung vượt quá giới hạn {MAX_DOCX_CONTENT_CHARS} ký tự."
        )

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run("BÁO CÁO PHÂN TÍCH PHÁP LÝ")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    generated_at = doc.add_paragraph()
    generated_at.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated_run = generated_at.add_run(
        f"Thời điểm tạo: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    generated_run.italic = True
    generated_run.font.size = Pt(10)

    disclaimer = doc.add_paragraph()
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    disclaimer_run = disclaimer.add_run(DISCLAIMER)
    disclaimer_run.italic = True
    disclaimer_run.font.size = Pt(10)

    doc.add_paragraph()

    for raw_line in content.splitlines():
        line = raw_line.strip()

        if not line:
            continue

        clean_line = line.lstrip("#").strip()

        if line.startswith("#") or clean_line.startswith(
            ("I.", "II.", "III.", "IV.", "V.", "VI.")
        ):
            paragraph = doc.add_paragraph()

            run = paragraph.add_run(clean_line)
            run.bold = True
            run.font.size = Pt(13)

            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)

        elif line.startswith(("- ", "* ", "• ")):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:].strip())
            paragraph.paragraph_format.space_after = Pt(3)

        else:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.add_run(clean_line)
            paragraph.paragraph_format.space_after = Pt(6)

    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)

    filename = (
        "Bao-cao-phan-tich-phap-ly-"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    )

    return stream, filename
